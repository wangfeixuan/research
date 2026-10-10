"""Create a short editable presentation from the measured Week 5 results."""
from pathlib import Path
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE

root=Path(__file__).resolve().parent
results=root/'results'
with (results/'representation_comparison.csv').open() as f: reps=list(csv.DictReader(f))
with (results/'vae_dimension_comparison.csv').open() as f: dims=list(csv.DictReader(f))
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY=RGBColor(30,48,69); BLUE=RGBColor(27,105,170); GRAY=RGBColor(68,83,99); LIGHT=RGBColor(226,237,248)
font='Arial'

def text(slide,x,y,w,h,s,size=22,bold=False,color=GRAY):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    tf.text=s
    for p in tf.paragraphs:
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
    return box

def base(title,number):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor(255,255,255)
    text(sl,.65,.35,12,0.6,title,30,True,NAVY)
    line=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.65),Inches(1.05),Inches(12.0),Inches(.025));line.fill.solid();line.fill.fore_color.rgb=LIGHT;line.line.fill.background()
    text(sl,12.0,7.10,.6,.25,str(number),11,False,GRAY)
    return sl

def bullet(sl,x,y,w,lines,size=21,gap=.55):
    for i,line in enumerate(lines): text(sl,x,y+i*gap,w,gap-.06,'•  '+line,size,False,GRAY)

def add_picture(sl,path,x,y,w,h): sl.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))

def table(sl,headers,rows,x,y,widths,row_h=.52):
    tb=sl.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(sum(widths)),Inches((len(rows)+1)*row_h)).table
    for j,w in enumerate(widths): tb.columns[j].width=Inches(w)
    for i,row in enumerate([headers]+rows):
        for j,value in enumerate(row):
            cell=tb.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.1);cell.margin_right=Inches(.07);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=LIGHT if i==0 else RGBColor(255,255,255)
            for p in cell.text_frame.paragraphs:
                p.font.name=font;p.font.size=Pt(16 if i==0 else 15);p.font.bold=i==0;p.font.color.rgb=NAVY
    return tb

sl=base('Learning Image Representations for Clustering',1)
text(sl,.75,1.55,11.6,1.0,'Fashion-MNIST: raw pixels, PCA, AE, and VAE',28,False,BLUE)
text(sl,.75,3.3,10,1.0,'Week 5 mini research project',23,False,GRAY)
text(sl,.75,5.5,11,0.8,'Question: how does the representation change K-means clustering?',24,True,NAVY)

sl=base('Question and dataset',2)
bullet(sl,.8,1.55,11.9,['Fashion-MNIST contains 28 × 28 grayscale clothing images.','I trained on 20,000 images and evaluated clustering on 10,000 test images.','Labels were used only after clustering to calculate NMI and ARI.'],22,.85)

sl=base('Four representations, one clustering method',3)
bullet(sl,.8,1.5,12,['Raw pixels: flatten each image to 784 values.','PCA: reduce pixel vectors to 20 components.','AE: use the encoder’s 20-value latent vector.','VAE: use the encoder’s 20-value mean vector.','K-means stays fixed at K = 10 for every comparison.'],20,.85)

sl=base('AE and VAE',4)
text(sl,.8,1.55,5.6,.5,'Autoencoder',25,True,BLUE)
text(sl,.8,2.15,5.6,2.2,'Image → encoder → z → decoder → reconstruction\n\nLearns to rebuild the input image.',21)
text(sl,6.8,1.55,5.6,.5,'Variational autoencoder',25,True,BLUE)
text(sl,6.8,2.15,5.6,2.6,'Image → encoder → μ and log σ² → sampled z → decoder\n\nReconstruction and KL loss shape the latent space.',21)
text(sl,.8,5.55,11.8,.8,'For VAE clustering, I used μ because it does not change with sampling.',22,True,NAVY)

sl=base('Experimental setup',5)
table(sl,['Setting','Value'],[['AE / VAE architecture','784 → 256 → 64 → latent → 64 → 256 → 784'],['Optimizer','Adam, learning rate 0.001'],['Batch size / epochs','256 / 5'],['Latent dimensions','2, 10, 20, 50'],['K-means','K = 10, n_init = 10, random_state = 42']],.8,1.5,[3.5,8.7],.7)
text(sl,.8,6.2,11.8,.55,'Week 3 classification: MLP 88.04%  |  CNN 91.55% (Fashion-MNIST test accuracy)',19,True,NAVY)

sl=base('Clustering comparison',6)
table(sl,['Representation','Dimension','NMI','ARI'],[[r['Representation'],r['Dimension'],f"{float(r['NMI']):.4f}",f"{float(r['ARI']):.4f}"] for r in reps],.8,1.5,[4.5,2.1,2.1,2.1],.68)
text(sl,.8,5.7,11.7,.9,'VAE-20 led this run; AE-20 fell below raw pixels and PCA-20.',22,True,NAVY)

sl=base('VAE dimension changes the result',7)
table(sl,['Dimension','Recon BCE','VAE loss','NMI','ARI'],[[r['latent_dim'],f"{float(r['test_mean_recon_bce_per_image']):.2f}",f"{float(r['test_vae_loss_per_image']):.2f}",f"{float(r['NMI']):.4f}",f"{float(r['ARI']):.4f}"] for r in dims],.8,1.5,[2.1,2.4,2.3,1.7,1.7],.63)
text(sl,.8,5.55,11.6,1.1,'Lowest reconstruction BCE and VAE loss: d = 10. Best clustering: d = 20.\nA lower VAE loss did not guarantee a higher NMI or ARI.',21,True,NAVY)

sl=base('Reconstruction and generation',8)
add_picture(sl,results/'vae_reconstructions_10.png',.7,1.45,11.9,2.0)
add_picture(sl,results/'vae_generated_20.png',4.65,3.6,4.1,3.2)
text(sl,.75,3.8,3.4,2.4,'Top: 10 inputs and their reconstructions.\n\nRight: 20 new images sampled from N(0, I).',18)
text(sl,9.0,4.1,3.4,1.7,'Many generated images are recognizable but blurry.',19,True,NAVY)

sl=base('Latent space and interpolation',9)
add_picture(sl,results/'vae_latent_2.png',.7,1.35,5.6,4.9)
add_picture(sl,results/'vae_interpolation.png',6.5,2.25,5.9,1.0)
text(sl,6.5,3.6,5.8,2.3,'The 2D latent space has some visible class regions, but many colors overlap.\n\nThe sneaker-to-T-shirt interpolation changes gradually.',20)

sl=base('Conclusions and next steps',10)
bullet(sl,.8,1.5,11.9,['Representation choice changed the K-means scores even though K-means stayed fixed.','VAE-20 reached NMI 0.5713 and ARI 0.4038 in this run.','Reconstruction quality alone did not predict clustering quality.','Next: repeat across seeds, train longer, and try a convolutional VAE.'],20,.95)
text(sl,.8,6.1,11.9,.7,'Main limit: five epochs and one random seed make this an initial result.',20,True,NAVY)

out=root/'week5_presentation.pptx';prs.save(out);print(out,len(prs.slides))
